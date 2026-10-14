import contextlib
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import armor_export as ae  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

SECRET = "cxup_TESTSECRET"
BN, PFX, DAY = "169699783343081", "Ty6aI0lN", "2026-10-11"


def mkrow(i, ts):
    return {"timestamp": ts, "applicationName": "armor", "api": "/resources/v1/review/businessId/*",
            "armor_request_id": f"rid{i}", "client_ip": "1.2.3.4", "full_uri": f"https://x/{i}",
            "method": "GET", "status": "429" if i % 5 == 0 else "200", "time_ms": 10 + i,
            "trace_id": f"t{i}"}


class Resp:
    def __init__(self, rows, extra=(), status=200):
        self.status_code = status
        lines = [json.dumps({"queryId": {"queryId": "q"}}),
                 json.dumps({"result": {"results": [{"userData": json.dumps(r)} for r in rows]}})]
        self.text = "\n".join(lines + [json.dumps(e) for e in extra])


def run_main(post, out):
    buf = io.StringIO()
    with mock.patch.dict(os.environ, {"CORALOGIX_API_KEY": SECRET}), \
            mock.patch.object(ae.requests, "post", post), contextlib.redirect_stdout(buf):
        ae.main(["--business-number", BN, "--key-prefix", PFX, "--day", DAY, "--out", out])
    return buf.getvalue()


class ArmorExportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.out = str(Path(self.tmp.name) / "o" / "x.xlsx")

    def tearDown(self):
        self.tmp.cleanup()

    def test_query_template(self):
        q = ae.build_query(BN, PFX)
        self.assertIn(f'x-business-number:"{BN}"', q)
        self.assertIn(f"contains('{PFX}')", q)
        self.assertNotIn("{{", q)
        self.assertTrue(q.rstrip().endswith("| limit 10000"))

    def test_rejects_unsafe_inputs(self):
        for bn, pfx in [("12ab", PFX), (BN, "x') || true //"), (BN, "")]:
            with self.assertRaises(SystemExit):
                ae.build_query(bn, pfx)

    def test_single_run_sorted_xlsx(self):
        rows = [mkrow(i, f"2026-10-11 {i % 24:02d}:00:00.000") for i in range(50)][::-1]
        calls = []

        def post(url, headers=None, json=None, timeout=None):
            calls.append(json["metadata"])
            return Resp(rows)

        out = run_main(post, self.out)
        self.assertEqual(len(calls), 1)
        self.assertEqual((calls[0]["startDate"], calls[0]["endDate"]),
                         ("2026-10-11T00:00:00Z", "2026-10-12T00:00:00Z"))
        self.assertEqual(calls[0]["tier"], "TIER_ARCHIVE")
        self.assertIn("IST: 11/10/2026 05:30 -> 12/10/2026 05:30", out)
        self.assertNotIn(SECRET, out)
        ws = load_workbook(self.out).active
        self.assertEqual([c.value for c in ws[1]], ae.COLUMNS)
        self.assertEqual((ws.max_row, ws.freeze_panes, ws.auto_filter.ref), (51, "A2", "A1:J51"))
        ts = [ws.cell(r, 1).value for r in range(2, 52)]
        self.assertEqual(ts, sorted(ts))
        self.assertIsInstance(ws["H2"].value, int)

    def test_scan_limit_splits_and_dedupes(self):
        calls = []

        def post(url, headers=None, json=None, timeout=None):
            md = json["metadata"]
            calls.append((md["startDate"], md["endDate"]))
            if len(calls) == 1:
                return Resp([], extra=[{"warning": {"bytesScannedLimitWarning": {}}}])
            h = int(md["startDate"][11:13])
            part = [mkrow(h * 100 + k, f"2026-10-11 {h:02d}:00:0{k}.000") for k in range(3)]
            if h == 6:
                part.append(mkrow(1, "2026-10-11 00:00:01.000"))  # boundary dup of a window-0 row
            return Resp(part)

        out = run_main(post, self.out)
        self.assertEqual(calls[1:], [("2026-10-11T00:00:00Z", "2026-10-11T06:00:00Z"),
                                     ("2026-10-11T06:00:00Z", "2026-10-11T12:00:00Z"),
                                     ("2026-10-11T12:00:00Z", "2026-10-11T18:00:00Z"),
                                     ("2026-10-11T18:00:00Z", "2026-10-12T00:00:00Z")])
        self.assertIn("merged 13 -> 12", out)
        self.assertIn("WARNING/ERROR LINE", out)
        ws = load_workbook(self.out).active
        ids = [ws.cell(r, 4).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(len(ids), len(set(ids)))

    def test_row_limit_triggers_split(self):
        n = []

        def post(url, headers=None, json=None, timeout=None):
            n.append(1)
            return Resp([mkrow(i, "2026-10-11 00:00:00.000") for i in range(ae.ROW_LIMIT)]
                        if len(n) == 1 else [])

        run_main(post, self.out)
        self.assertEqual(len(n), 5)

    def test_http_error_exits_without_key(self):
        with self.assertRaises(SystemExit) as cm:
            run_main(lambda *a, **k: Resp([], status=403), self.out)
        self.assertNotIn(SECRET, str(cm.exception))

    def test_missing_key_exits(self):
        with mock.patch.dict(os.environ, {}, clear=True), self.assertRaises(SystemExit):
            ae.load_key()


if __name__ == "__main__":
    unittest.main()
